import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

SCREENSHOT_DIR = r"C:\Users\Acer\.gemini\antigravity-ide\brain\a67442fc-448d-4dd9-82c2-2f47df20c631\screenshots"
LOGO_PATH = r"c:\Users\Acer\projects\ERP\backend\assets\logo.png"
OUTPUT_DOCX = r"c:\Users\Acer\projects\ERP\College_ERP_Demo_Walkthrough.docx"
ARTIFACT_DOCX = r"C:\Users\Acer\.gemini\antigravity-ide\brain\a67442fc-448d-4dd9-82c2-2f47df20c631\College_ERP_Demo_Walkthrough.docx"

# Color Palette
COLOR_PRIMARY = RGBColor(30, 58, 138)     # Navy #1E3A8A
COLOR_SECONDARY = RGBColor(15, 118, 110)  # Teal #0F766E
COLOR_DARK = RGBColor(15, 23, 42)         # Slate 900 #0F172A
COLOR_MUTED = RGBColor(71, 85, 105)       # Slate 600 #475569
HEX_LIGHT_BG = "F8FAFC"                   # Light gray/slate
HEX_BORDER = "CBD5E1"                     # Slate 300
HEX_THEAD = "1E3A8A"                      # Navy
HEX_ACCENT_BG = "EFF6FF"                  # Light blue

def set_cell_shading(cell, color_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_header(doc, text, level=1, color=COLOR_PRIMARY):
    p = doc.add_heading(level=level)
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.color.rgb = color
    run.bold = True
    if level == 1:
        run.font.size = Pt(18)
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
    elif level == 2:
        run.font.size = Pt(14)
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
    elif level == 3:
        run.font.size = Pt(12)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
    return p

def add_callout(doc, text, title=""):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, HEX_ACCENT_BG)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    
    if title:
        run_t = p.add_run(f"📌 {title}\n")
        run_t.bold = True
        run_t.font.name = 'Calibri'
        run_t.font.size = Pt(10.5)
        run_t.font.color.rgb = COLOR_PRIMARY
        
    run_body = p.add_run(text)
    run_body.font.name = 'Calibri'
    run_body.font.size = Pt(10)
    run_body.font.color.rgb = COLOR_DARK
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_feature_section(doc, title, screenshot_filename, purpose_text, buttons_list):
    """
    Renders one clean subsection:
    1. Subsection Heading
    2. Purpose paragraph
    3. Screenshot image (6.2 inches wide)
    4. Button breakdown table
    """
    add_header(doc, title, level=2, color=COLOR_PRIMARY)
    
    # Purpose
    p_purpose = doc.add_paragraph()
    p_purpose.paragraph_format.space_after = Pt(6)
    run_p_label = p_purpose.add_run("Screen Purpose: ")
    run_p_label.bold = True
    run_p_label.font.name = 'Calibri'
    run_p_label.font.size = Pt(10.5)
    run_p_label.font.color.rgb = COLOR_SECONDARY
    
    run_p_text = p_purpose.add_run(purpose_text)
    run_p_text.font.name = 'Calibri'
    run_p_text.font.size = Pt(10.5)
    run_p_text.font.color.rgb = COLOR_DARK
    
    # Screenshot Image
    img_path = os.path.join(SCREENSHOT_DIR, screenshot_filename)
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(6.2))
        
        # Caption
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(8)
        run_cap = p_cap.add_run(f"Figure: Live demonstration view of {title}")
        run_cap.italic = True
        run_cap.font.name = 'Calibri'
        run_cap.font.size = Pt(9)
        run_cap.font.color.rgb = COLOR_MUTED
    else:
        add_callout(doc, f"Image file {screenshot_filename} not found.", "Notice")

    # Button Breakdown Table
    if buttons_list:
        p_btn_title = doc.add_paragraph()
        p_btn_title.paragraph_format.space_before = Pt(6)
        p_btn_title.paragraph_format.space_after = Pt(2)
        run_btn_hdr = p_btn_title.add_run("Buttons & Controls on this Screen:")
        run_btn_hdr.bold = True
        run_btn_hdr.font.name = 'Calibri'
        run_btn_hdr.font.size = Pt(10)
        run_btn_hdr.font.color.rgb = COLOR_PRIMARY
        
        table = doc.add_table(rows=len(buttons_list) + 1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        
        # Headers
        headers = ["Button / Action Element", "What Happens When You Click It"]
        col_widths = [Inches(2.2), Inches(4.3)]
        
        for col_idx, (text, width) in enumerate(zip(headers, col_widths)):
            cell = table.cell(0, col_idx)
            cell.width = width
            set_cell_shading(cell, HEX_THEAD)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(255, 255, 255)
            
        # Rows
        for row_idx, (btn_name, btn_desc) in enumerate(buttons_list, start=1):
            bg_color = HEX_LIGHT_BG if row_idx % 2 == 0 else "FFFFFF"
            for col_idx, (val, width) in enumerate(zip([btn_name, btn_desc], col_widths)):
                cell = table.cell(row_idx, col_idx)
                cell.width = width
                set_cell_shading(cell, bg_color)
                set_cell_margins(cell, top=70, bottom=70, left=120, right=120)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(val)
                run.font.name = 'Calibri'
                run.font.size = Pt(9)
                if col_idx == 0:
                    run.bold = True
                    run.font.color.rgb = COLOR_PRIMARY
                else:
                    run.font.color.rgb = COLOR_DARK

        doc.add_paragraph().paragraph_format.space_after = Pt(10)

def build_walkthrough():
    print("[INFO] Creating College ERP Demo Walkthrough document...")
    doc = docx.Document()
    
    # Page setup - Standard Letter, 0.75 in margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # ----------------------------------------------------
    # COVER PAGE
    # ----------------------------------------------------
    p_logo = doc.add_paragraph()
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_before = Pt(40)
    p_logo.paragraph_format.space_after = Pt(20)
    if os.path.exists(LOGO_PATH):
        p_logo.add_run().add_picture(LOGO_PATH, width=Inches(1.8))

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("COLLEGE ERP SYSTEM")
    run_title.bold = True
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(26)
    run_title.font.color.rgb = COLOR_PRIMARY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(12)
    run_sub = p_sub.add_run("Working Demonstration Walkthrough & Feature Guide")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(15)
    run_sub.font.color.rgb = COLOR_SECONDARY

    p_desc = doc.add_paragraph()
    p_desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_desc.paragraph_format.space_after = Pt(40)
    run_desc = p_desc.add_run("A visual, plain-English overview of every screen, button, and action\ndesigned for institutional review and administration feedback.")
    run_desc.font.name = 'Calibri'
    run_desc.font.size = Pt(11)
    run_desc.font.color.rgb = COLOR_MUTED

    # Meta Table on Cover
    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_info = [
        ("Prepared For:", "College Principal, Administrator & Governing Council"),
        ("System Name:", "Aura Institutional ERP (Admissions, Fees, Attendance & Library)"),
        ("Review Version:", "Working Interactive Demonstration (Phase 1 Review)"),
        ("Date of Review:", "September 2026")
    ]
    for i, (k, v) in enumerate(meta_info):
        c0 = table.cell(i, 0)
        c1 = table.cell(i, 1)
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        set_cell_margins(c0, top=60, bottom=60, left=100, right=100)
        set_cell_margins(c1, top=60, bottom=60, left=100, right=100)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.bold = True
        r0.font.name = 'Calibri'
        r0.font.size = Pt(10)
        r0.font.color.rgb = COLOR_PRIMARY
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        r1.font.name = 'Calibri'
        r1.font.size = Pt(10)
        r1.font.color.rgb = COLOR_DARK

    doc.add_page_break()

    # ----------------------------------------------------
    # 1. INTRODUCTION
    # ----------------------------------------------------
    add_header(doc, "1. Introduction", level=1)
    
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_after = Pt(8)
    p_intro.paragraph_format.line_spacing = 1.15
    run_intro = p_intro.add_run(
        "Welcome to the visual walkthrough of your new College ERP system. "
        "This platform is designed to bring your entire institution together into one clean digital environment, "
        "replacing handwritten registers, paper fee receipts, and delayed attendance sheets with instant, automatic records. "
        "This document presents a live, working demonstration created specifically for your evaluation. "
        "It is not a final, locked-in product — we have prepared it so you can review each screen, test the workflow, "
        "and tell us exactly what you want adjusted to suit your college's daily routines."
    )
    run_intro.font.name = 'Calibri'
    run_intro.font.size = Pt(11)
    run_intro.font.color.rgb = COLOR_DARK

    add_callout(
        doc,
        "How to use this document: As you read through the screens below, you can point to any button, table, or box "
        "and note your feedback using the simple questions provided at the very end of this guide.",
        "A Note for the Administrator"
    )

    # ----------------------------------------------------
    # HOW ACCESS & LOGIN WORKS
    # ----------------------------------------------------
    add_header(doc, "How Staff & Students Access the System", level=2)
    
    p_login_desc = doc.add_paragraph()
    p_login_desc.paragraph_format.space_after = Pt(6)
    p_login_desc.paragraph_format.line_spacing = 1.15
    r_ld = p_login_desc.add_run(
        "Every person who uses the system visits the same single sign-in page on their computer or mobile phone. "
        "The system instantly recognizes who they are based on their email and shows them only the tools they are allowed to see: "
        "an Administrator sees the whole college; a Clerk sees fee counters; a Teacher sees their class roll; and a Student sees only their own fees and attendance."
    )
    r_ld.font.name = 'Calibri'
    r_ld.font.size = Pt(10.5)
    r_ld.font.color.rgb = COLOR_DARK

    add_feature_section(
        doc,
        "Unified Sign-In Screen",
        "01_login_screen.png",
        "This is the front door of your college portal. Staff and students enter their registered institutional email and secure password here.",
        [
            ("Email Address / User ID", "A typing box where the staff member or student types their college email."),
            ("Password", "A private box where the user types their password (letters are hidden with dots for safety)."),
            ("Sign In Button", "Pressing this checks the password and takes the user directly into their own dashboard."),
            ("Quick Demo Access: Admin", "One-click shortcut for testing that fills in the Administrator account automatically."),
            ("Quick Demo Access: Teacher", "One-click shortcut for testing that fills in a Professor account automatically."),
            ("Quick Demo Access: Clerk", "One-click shortcut for testing that fills in the Cashier/Clerk account automatically."),
            ("Quick Demo Access: Student", "One-click shortcut for testing that fills in a Student account automatically.")
        ]
    )

    doc.add_page_break()

    # ----------------------------------------------------
    # 2. ADMINISTRATOR PORTAL
    # ----------------------------------------------------
    add_header(doc, "2. Administrator Portal", level=1)
    
    p_admin_intro = doc.add_paragraph()
    p_admin_intro.paragraph_format.space_after = Pt(8)
    p_admin_intro.paragraph_format.line_spacing = 1.15
    r_ai = p_admin_intro.add_run(
        "The Administrator Portal is the master command center for the Principal, Director, and Office Superintendent. "
        "From here, leadership has a live 360-degree view of total fees collected, unpaid student dues, daily attendance percentages, "
        "library circulation, and faculty workloads."
    )
    r_ai.font.name = 'Calibri'
    r_ai.font.size = Pt(10.5)
    r_ai.font.color.rgb = COLOR_DARK

    # Screen 2: Admin Dashboard
    add_feature_section(
        doc,
        "Executive Overview Dashboard",
        "02_admin_dashboard.png",
        "This is the first screen leadership sees upon signing in. It summarizes total fee billing, collections received to date, outstanding student balances, and banking reversals.",
        [
            ("Total Assigned Fees Card", "Shows the total monetary value of all fees billed across all enrolled students this year."),
            ("Total Revenue Collected Card", "Shows the actual cash, cheques, and online payments deposited into the college account so far."),
            ("Outstanding Pending Dues Card", "Shows the remaining money still to be collected from students who have not finished paying."),
            ("Audited Reversals Card", "Shows the total value of bounced cheques or cancelled fee receipts that were safely undone."),
            ("Detailed Reports Button", "Jumps directly to printable financial spreadsheets and banking breakdown sheets."),
            ("Manage Fees Button", "Opens the fee structure tool where you set tuition amounts for new academic batches."),
            ("Revenue Collection Timeline", "A line graph displaying the pace of fee collections day-by-day across this month."),
            ("Fee Head Split Chart", "A visual pie graph showing what percentage of income came from Tuition, Hostel, Exam, or Transport."),
            ("Program-wise Assigned vs Collected", "A bar chart comparing collection performance between B.Tech, BBA, and MBA programs.")
        ]
    )

    # Screen 3: Fee Structures
    add_feature_section(
        doc,
        "Fee Structure & Pricing Rules",
        "03_admin_fee_structures.png",
        "This screen lists the approved fees for every branch and batch. Here you define how much students in each program owe for Tuition, Hostel, Library, Lab, and Transport.",
        [
            ("New Fee Structure Button", "Opens a simple form to create a new fee rule (for example, introducing an MCA or Diploma fee)."),
            ("Academic Year Filter", "Lets you switch between the current academic year and previous years to review historical pricing."),
            ("Class & Branch Selector", "Filters the list so you can view only B.Tech, BBA, or MBA fees at a time."),
            ("Fee Amount Column", "Displays the official approved rate in Indian Rupees for that particular item."),
            ("Due Date Column", "Shows the deadline date before which students must deposit this fee."),
            ("Edit Icon (Pencil)", "Lets you update the amount or extend the due date for that specific fee item."),
            ("Delete Icon (Trash)", "Deletes a fee rule that was entered by accident (only allowed if no student has paid against it yet).")
        ]
    )

    # Screen 4: Students Directory
    add_feature_section(
        doc,
        "Student Records & Admissions Directory",
        "04_admin_students_list.png",
        "The master institutional roll containing every enrolled student with their photo, roll number, class section, guardian phone number, and real-time fee balance status.",
        [
            ("New Admission Button", "Opens the official student registration form to enroll a newly admitted student."),
            ("Search Input Box", "Type any part of a student's name, roll number, or admission number to instantly find them."),
            ("Class & Section Filter", "Filters the list to view one specific classroom (such as B.Tech Computer Science Section A)."),
            ("Status Filter", "Lets you toggle between Active students, Graduated students, or Deactivated students."),
            ("Student Photo Avatar", "Displays the student's passport photo for quick visual identity verification."),
            ("Student Name Link", "Clicking the student's name opens their complete 360-degree personal academic and fee file."),
            ("Outstanding Dues Badge", "Color-coded tag showing whether their fees are fully Cleared (green), Partial (orange), or Unpaid (red)."),
            ("View Profile Icon (Eye)", "Opens the student's detailed academic file with guardian details and full attendance record."),
            ("Edit Record Icon (Pencil)", "Allows modifying contact phone numbers, home address, or guardian details."),
            ("Assign Fees Icon (Layers)", "Manually attaches or detaches specific optional fee heads (like adding Hostel or Bus facility)."),
            ("Deactivate Icon (User X)", "Safely marks a student as transferred or withdrawn without deleting their financial history.")
        ]
    )

    # Screen 5: New Admission Form
    add_feature_section(
        doc,
        "New Student Admission Form",
        "05_admin_new_student_modal.png",
        "The digital enrollment form used when a new student joins the college. It captures personal information, parent contact numbers, and submitted certificates.",
        [
            ("Full Legal Name", "Box to record the student's official name as printed on their 10th grade mark sheet."),
            ("College Email Address", "Generates the student's official email address used for institutional notices and portal sign-in."),
            ("Roll Number & Admission No.", "Assigns the unique institutional identification code used on all exams and receipts."),
            ("Program & Class Dropdown", "Selects the admitted branch (e.g. B.Tech Computer Engineering, BBA, MBA)."),
            ("Section Selector", "Places the student in Section A, B, or C for classroom roll calls."),
            ("Admission Date", "Calendar picker to record the official date fee was paid and admission confirmed."),
            ("Guardian Name & Relationship", "Records parent/guardian details (Father, Mother, or Guardian) for emergency contact."),
            ("Guardian Mobile Number", "Phone number stored for attendance alerts, fee notices, and urgent communications."),
            ("Document Checklist (Aadhaar, TC, Mark sheets)", "Checkboxes where the admissions officer ticks off each original certificate submitted."),
            ("Submit Admission Button", "Saves the student profile and automatically generates their semester fee dues in the system."),
            ("Cancel Button", "Closes the admission window without saving any information.")
        ]
    )

    # Screen 6: Student 360 Profile
    add_feature_section(
        doc,
        "Unified Student 360 Profile",
        "06_admin_student_profile.png",
        "A complete single-page dossier for any student. Shows their family contact info, verification status of documents, attendance percentage, and their entire payment history.",
        [
            ("Profile Header", "Displays student photo, roll number, admitted course, active status tag, and email."),
            ("Guardian & Emergency Tile", "Shows parent name, phone number, relationship, and permanent home address."),
            ("Document Verification Badges", "Green checkmarks showing verified Aadhaar, Transfer Certificate, and Mark sheets on file."),
            ("Fee Ledger Overview", "Cards displaying total fees assigned, total paid, and balance remaining for the current year."),
            ("Fee Particulars Breakdown", "Table showing exact progress on Tuition, Lab, Library, and Exam fees separately."),
            ("Official Receipts Table", "Lists every receipt voucher issued to this student with date, payment mode, and download link."),
            ("Download Voucher Button", "Generates a clean PDF copy of any historical fee receipt for the student or parent.")
        ]
    )

    # Screen 7: Teachers Management
    add_feature_section(
        doc,
        "Faculty & Teacher Directory",
        "07_admin_teachers_list.png",
        "The directory of all teaching faculty members. Shows the subjects each professor teaches, their employee ID, contact details, and their assigned classroom sections.",
        [
            ("Add Faculty Member Button", "Opens a dialog to register a new lecturer or professor in the system."),
            ("Department Filter", "Filters faculty by department (Computer Science, Information Technology, Management)."),
            ("Teacher Name & Photo", "Shows professor's name and institutional email address."),
            ("Employee Identification Tag", "Unique staff ID used for classroom scheduling and attendance audit."),
            ("Subjects Taught Badges", "Displays the specific course papers the teacher instructs (e.g. Operating Systems, Algorithms)."),
            ("Class Assigned Tag", "Shows the specific classroom sections where this teacher has permission to take attendance."),
            ("Edit Faculty Details", "Opens an update window to assign new subjects or change assigned classroom sections.")
        ]
    )

    # Screen 8: Attendance Analytics
    add_feature_section(
        doc,
        "Institutional Attendance & Defaulters Report",
        "08_admin_attendance_analytics.png",
        "A bird's-eye view of classroom attendance. Automatically highlights students falling below the mandatory 75% attendance threshold so notices can be sent before exams.",
        [
            ("Date Range Picker", "Allows analyzing attendance over any custom period (last week, last month, or full semester)."),
            ("Class & Section Filter", "Selects a specific classroom to inspect its daily attendance records."),
            ("Overall College Attendance Rate", "Big-number metric showing the average attendance across the entire institution today."),
            ("Defaulter Warning List", "A dedicated table listing students under 75% attendance with their exact percentage in red."),
            ("Export Attendance Register", "Downloads a spreadsheet register suitable for university exam eligibility audits.")
        ]
    )

    # Screen 9: Library Catalog
    add_feature_section(
        doc,
        "Central Library Catalog & Inventory",
        "09_admin_library_catalog.png",
        "The inventory register of all physical books in the college library. Tracks total copies owned, copies currently on loan to students, and copies on shelf.",
        [
            ("Add New Book Button", "Opens a form to catalog a newly purchased textbook with title, author, ISBN, and quantity."),
            ("Category Filter", "Filters books by subject area (Computer Science, Finance, Mathematics, Electronics, Physics)."),
            ("Search Book Title / Author", "Finds any textbook instantly by typing part of the title or author name."),
            ("Total Copies Column", "Shows the total number of physical copies owned by the college."),
            ("Available Shelf Copies", "Real-time count of how many copies are currently available on the shelf to be borrowed."),
            ("Circulation Status Badge", "Displays 'In Stock' if copies are available, or 'Fully Borrowed' if all are out on loan.")
        ]
    )

    # Screen 10: Clerks & Cashiers
    add_feature_section(
        doc,
        "Cashier & Staff User Accounts",
        "10_admin_clerks_cashiers.png",
        "Management screen for accounts staff, cashiers, and bursars. Allows creating cashier logins, assigning counter permissions, and auditing staff activity.",
        [
            ("Register New Staff User", "Creates a new clerk account with designated login email and counter password."),
            ("Staff Directory Table", "Lists all active counter clerks with their name, work email, and last login time."),
            ("Reset Password Button", "Allows the administrator to issue a temporary password if a staff member forgets theirs."),
            ("Suspend / Activate Account", "One-click switch to lock a cashier's access if they leave the college or go on leave.")
        ]
    )

    # Screen 11: Financial Reports
    add_feature_section(
        doc,
        "Financial Reports & Collection Analytics",
        "11_admin_financial_reports.png",
        "High-level accounting and bursar summaries. Breaks down collections by payment mode (Cash, Bank Cheques, DD, and Online), program, and date.",
        [
            ("Collections by Payment Mode", "Shows exact amounts received in Cash, Bank Cheque, Demand Draft, and Online."),
            ("Branch Recovery Comparison", "Compares expected fee recovery vs. collected amounts across B.Tech, BBA, and MBA."),
            ("Daily Cash Settlement Register", "Displays each cashier's total cash collected today for daily bank deposits."),
            ("Print Summary Button", "Formats the financial report onto a clean printable sheet for the Board of Trustees meeting."),
            ("Download Excel Spreadsheet", "Exports the raw accounting rows for your external chartered accountant.")
        ]
    )

    # Screen 12: Settings
    add_feature_section(
        doc,
        "Institution Profile & System Settings",
        "12_admin_settings.png",
        "General configuration for the college. Stores the official college name, address, receipt numbering format, and academic calendar dates.",
        [
            ("College Legal Name", "Box to set the official institutional title printed at the top of every fee receipt."),
            ("Campus Address & Phone", "Contact details and affiliation number displayed on student fee vouchers."),
            ("Receipt Voucher Prefix", "Allows choosing how receipts are numbered (e.g., 'RCP-2026-00001')."),
            ("Academic Session Selector", "Sets the active academic year for attendance and fee calculations (e.g., '2026-2027')."),
            ("Save Settings Button", "Applies and locks all updated institutional details across the entire platform.")
        ]
    )

    doc.add_page_break()

    # ----------------------------------------------------
    # 3. CLERK & CASHIER PORTAL
    # ----------------------------------------------------
    add_header(doc, "3. Clerk & Accounts Counter Portal", level=1)
    
    p_clerk_intro = doc.add_paragraph()
    p_clerk_intro.paragraph_format.space_after = Pt(8)
    p_clerk_intro.paragraph_format.line_spacing = 1.15
    r_ci = p_clerk_intro.add_run(
        "The Clerk Portal is designed for front-desk cashiers and accounts staff. "
        "Its layout is optimized for rapid student lookups, entering fee payments across different modes, "
        "generating verified PDF receipts with college logos, and managing textbook loans at the library circulation desk."
    )
    r_ci.font.name = 'Calibri'
    r_ci.font.size = Pt(10.5)
    r_ci.font.color.rgb = COLOR_DARK

    # Screen 13: Clerk Dashboard
    add_feature_section(
        doc,
        "Cashier Station Overview",
        "13_clerk_dashboard.png",
        "The opening dashboard for the cashier. Shows today's cash collections, total receipts issued during this shift, and quick-action buttons for common tasks.",
        [
            ("Today's Collections Tile", "Shows the total money collected by this specific cashier since logging in today."),
            ("Vouchers Issued Tile", "Count of how many official fee receipts this workstation printed today."),
            ("Quick Fee Collection Button", "Shortcut that opens the student fee collection counter immediately."),
            ("Library Circulation Button", "Shortcut to the book issue and return counter."),
            ("Recent Counter Transactions", "Table showing the last few deposits taken with student roll numbers and amounts.")
        ]
    )

    # Screen 14: Collect Fees Counter
    add_feature_section(
        doc,
        "Fee Collection & Receipt Counter",
        "14_clerk_collect_fees.png",
        "The primary screen where students pay their semester dues. Typing a roll number displays the student's name, their unpaid fee items, and payment options.",
        [
            ("Student Search Box", "Type the student's roll number or name (e.g. 'Priya') to instantly load their billing account."),
            ("Student Particulars Card", "Displays student photo, roll number, class, and guardian phone number for positive identity verification."),
            ("Fee Head Selector", "Choose which specific fee is being paid (Tuition Fee, Hostel Fee, or Exam Fee)."),
            ("Payment Amount Input", "Type the exact amount being deposited today (supports full payment or partial installment)."),
            ("Payment Mode Options", "Clickable buttons to select Cash, Cheque, Demand Draft, or Online Transfer."),
            ("Reference / Instrument No.", "Box to record Cheque number, Bank DD number, or Bank Transaction Reference."),
            ("Collect Fee & Issue Receipt", "Finalizes the payment, deducts the student's balance, and immediately opens a printable receipt.")
        ]
    )

    # Screen 15: Payment History
    add_feature_section(
        doc,
        "Payment History & Audit Ledger",
        "15_clerk_payment_history.png",
        "The chronological cashbook of all fee deposits. Every entry is logged with its permanent voucher number, student details, payment instrument, and clerk ID.",
        [
            ("Search by Receipt Number", "Instantly pulls up an old transaction by typing the number printed on the student's paper slip."),
            ("Search by Roll Number", "Displays all payments ever made by a single student throughout their college career."),
            ("Status Filter", "Filters between Successful receipts and Reversed receipts (bounced instruments)."),
            ("Download PDF Receipt Button", "Re-downloads an official PDF copy of the fee receipt for a student who lost their original slip."),
            ("Request Reversal Button", "Audited safety action to flag a bounced cheque and restore the unpaid balance on the student's ledger.")
        ]
    )

    # Screen 16: Pending Dues Register
    add_feature_section(
        doc,
        "Pending Dues & Defaulters Register",
        "16_clerk_pending_dues.png",
        "An organized list of students who have not completed their fee payments. Sorted by class and program so clerks can follow up systematically.",
        [
            ("Class / Program Filter", "Filter unpaid lists by B.Tech, BBA, or MBA to check branch-wise collection status."),
            ("Student Name & Roll No.", "Shows the student and their parent contact number for quick phone reminders."),
            ("Total Fee Assigned", "Displays the total annual fees billed for this student."),
            ("Total Paid So Far", "Shows the amount already deposited in previous installments."),
            ("Pending Balance Due", "Prominently displays the remaining balance owed to the college in bold orange."),
            ("Direct Collect Shortcut", "Clicking the payment icon jumps directly into the collection screen with this student pre-loaded.")
        ]
    )

    # Screen 17: Library Circulation
    add_feature_section(
        doc,
        "Library Desk Circulation (Issue & Return)",
        "17_clerk_library_circulation.png",
        "The counter where library staff issue textbooks to students, process book returns, and collect overdue late fines.",
        [
            ("Issue Book Tab", "Simple form where the librarian selects a student, selects a book, and clicks to lend it out."),
            ("Active Loans Table", "Lists every book currently checked out, showing borrower name, book title, issue date, and due date."),
            ("Overdue Indicator", "Highlights books kept past the 14-day limit with calculated late fees (₹2 per day)."),
            ("Process Return Button", "Marks a book as returned to the shelf and updates the available copy count automatically."),
            ("Collect Fine Button", "Records payment of the late fine and marks the borrower's library account clear.")
        ]
    )

    doc.add_page_break()

    # ----------------------------------------------------
    # 4. TEACHER PORTAL
    # ----------------------------------------------------
    add_header(doc, "4. Teacher & Faculty Portal", level=1)
    
    p_teacher_intro = doc.add_paragraph()
    p_teacher_intro.paragraph_format.space_after = Pt(8)
    p_teacher_intro.paragraph_format.line_spacing = 1.15
    r_ti = p_teacher_intro.add_run(
        "The Teacher Portal provides professors and lecturers with a distraction-free mobile and desktop interface. "
        "Teachers can open their assigned classroom roll call, mark attendance in under 30 seconds with one tap, "
        "and review historical lecture logs."
    )
    r_ti.font.name = 'Calibri'
    r_ti.font.size = Pt(10.5)
    r_ti.font.color.rgb = COLOR_DARK

    # Screen 18: Teacher Dashboard
    add_feature_section(
        doc,
        "Faculty Overview & Schedule",
        "18_teacher_dashboard.png",
        "The opening screen for teaching staff. Shows the professor's assigned classes, today's lecture schedule, and whether attendance has been completed today.",
        [
            ("Assigned Class Badge", "Displays the teacher's designated branch (e.g. B.Tech CSE Section A)."),
            ("Today's Attendance Status", "Shows green if attendance is already completed today, or an orange alert if pending."),
            ("Quick Mark Attendance Button", "Launches the daily classroom roll call sheet immediately."),
            ("Total Students Enrolled", "Displays the roster size for the assigned class."),
            ("Recent Lecture Sessions", "Shows dates and attendance turnouts for the professor's last five class sessions.")
        ]
    )

    # Screen 19: Mark Attendance
    add_feature_section(
        doc,
        "Classroom Attendance Sheet",
        "19_teacher_mark_attendance.png",
        "The digital attendance register for taking roll in class. Pre-populates all students as 'Present' so the professor only taps the few students who are absent.",
        [
            ("Class & Section Selector", "Choose which class period you are taking attendance for."),
            ("Lecture Date Picker", "Defaults to today's date, but allows selecting a previous date if entering records retrospectively."),
            ("Mark All Present Button", "Instantly sets every student on the list to Present with a single tap."),
            ("Present Status Pill (P)", "Turns green when the student is present in the classroom."),
            ("Absent Status Pill (A)", "Turns red when tapped, marking the student absent for this lecture."),
            ("Late Status Pill (L)", "Turns yellow when tapped, recording that the student arrived late."),
            ("Optional Remarks Box", "Allows the teacher to type a note (e.g. 'Excused for sports meet' or 'Sick leave')."),
            ("Save Attendance Record", "Locks the attendance session, updates college registers, and flags any new defaulters.")
        ]
    )

    # Screen 20: Teacher History
    add_feature_section(
        doc,
        "Teacher Lecture Log & History",
        "20_teacher_attendance_history.png",
        "A historical archive of every lecture session conducted by this professor. Displays attendance rates and allows reviewing past student entries.",
        [
            ("Date Filter", "Select a month or week to review previous class lectures."),
            ("Lecture Session List", "Displays each conducted lecture date, total enrolled, total present, and overall percentage."),
            ("View Session Details (Eye)", "Opens the full roll call list for that past date to see who was absent."),
            ("Edit Session Button", "Allows correcting an accidental mark within the institution's allowed revision window.")
        ]
    )

    doc.add_page_break()

    # ----------------------------------------------------
    # 5. STUDENT PORTAL
    # ----------------------------------------------------
    add_header(doc, "5. Student & Parent Portal", level=1)
    
    p_student_intro = doc.add_paragraph()
    p_student_intro.paragraph_format.space_after = Pt(8)
    p_student_intro.paragraph_format.line_spacing = 1.15
    r_si = p_student_intro.add_run(
        "The Student Portal gives students and their parents total transparency over their academic standing. "
        "Students can log in from home to check their attendance percentage, inspect fee dues, download official receipts, "
        "and track borrowed library books without needing to stand in long office lines."
    )
    r_si.font.name = 'Calibri'
    r_si.font.size = Pt(10.5)
    r_si.font.color.rgb = COLOR_DARK

    # Screen 21: Student Dashboard
    add_feature_section(
        doc,
        "Student Overview & Ledger",
        "21_student_dashboard.png",
        "The student's personal home screen. Highlights their attendance percentage dial, outstanding fee balance, semester fee breakdown, and transaction history.",
        [
            ("Attendance Percentage Ring", "Visual circular indicator showing the student's overall attendance rate (warning shown if below 75%)."),
            ("Outstanding Fee Due Tile", "Shows remaining unpaid balance in Rupees with the upcoming payment deadline date."),
            ("Semester Particulars Breakdown", "Cards displaying Tuition, Exam, Hostel, and Transport progress with visual percentage bars."),
            ("Pay Outstanding Fees Button", "Opens the payment dialog to clear pending dues."),
            ("Official Payment History Table", "Lists every receipt issued to the student with payment mode and voucher number."),
            ("Download Receipt Voucher", "Lets the student download a verified PDF receipt directly to their phone or laptop.")
        ]
    )

    # Screen 22: Pay Online Modal
    add_feature_section(
        doc,
        "Fee Payment Dialog (Current Demo View)",
        "22_student_pay_online_modal.png",
        "The window that opens when a student clicks 'Pay Outstanding Fees'. In this working demo, it shows clear instructions for paying at the accounts counter or bank.",
        [
            ("Due Amount Display", "Clearly shows the pending balance in Rupees that needs to be paid."),
            ("Student Roll Number Confirmation", "Confirms the student's registration ID so payments are credited correctly."),
            ("Accounts Office Instructions", "Simple instructions detailing where the student can deposit fees on campus."),
            ("Bank Transfer / NEFT Details", "College bank account number and IFSC code for direct wire transfers."),
            ("Understood / Close Button", "Dismisses the payment modal.")
        ]
    )

    # Screen 23: Student Attendance
    add_feature_section(
        doc,
        "Student Attendance Tracking",
        "23_student_attendance.png",
        "Detailed attendance ledger for the student. Shows their attendance standing across each subject and lists recent lecture dates with Present or Absent marks.",
        [
            ("Overall Attendance Percentage", "Displays cumulative semester attendance with clear eligibility color (Green: Safe, Red: Below 75%)."),
            ("Lectures Attended vs Conducted", "Shows exact counts (e.g. '42 lectures attended out of 50 conducted')."),
            ("Subject-Wise Breakdown", "Shows attendance separately for each enrolled course paper."),
            ("Daily Lecture Log", "Date-by-date list showing which days the student was marked Present, Absent, or Late.")
        ]
    )

    # Screen 24: Student Library
    add_feature_section(
        doc,
        "Student Borrowed Books & Returns",
        "24_student_library.png",
        "Personal library borrowing card. Shows all books currently checked out, scheduled return due dates, and any overdue fine balances.",
        [
            ("Currently Borrowed Count", "Shows how many textbooks the student currently has at home."),
            ("Book Title & Author", "Lists the title, author, and library catalog code of each borrowed book."),
            ("Return Due Date", "Displays the deadline date to return the book to avoid late charges."),
            ("Fine Status Tag", "Shows 'Active Loan - On Time' or flags late charges in red if the book is overdue.")
        ]
    )

    # Screen 25: Student Profile
    add_feature_section(
        doc,
        "Student Identity & Admissions Record",
        "25_student_profile.png",
        "The official institutional identity card of the student. Shows registered family contact details, admitted academic program, and certificate status.",
        [
            ("Student Identity Header", "Displays full legal name, official student photo, roll number, and admission number."),
            ("Admitted Course & Batch", "Shows degree program, academic year of admission, and assigned section."),
            ("Parent / Guardian Contact", "Lists father/mother's phone number and home address on official record."),
            ("Verified Documents Checklist", "Displays status of Aadhaar, Transfer Certificate, and previous mark sheets.")
        ]
    )

    doc.add_page_break()

    # ----------------------------------------------------
    # 6. WHAT IS INCOMPLETE / PLANNED FOR LATER
    # ----------------------------------------------------
    add_header(doc, "6. What is Incomplete or Planned for a Later Phase", level=1)
    
    p_inc_intro = doc.add_paragraph()
    p_inc_intro.paragraph_format.space_after = Pt(8)
    p_inc_intro.paragraph_format.line_spacing = 1.15
    r_ii = p_inc_intro.add_run(
        "To ensure complete clarity and honesty, this section outlines which features in this live demonstration are currently working "
        "as simulated placeholders and will be connected to external services in Phase 2 based on your approval."
    )
    r_ii.font.name = 'Calibri'
    r_ii.font.size = Pt(10.5)
    r_ii.font.color.rgb = COLOR_DARK

    table_inc = doc.add_table(rows=5, cols=3)
    table_inc.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_inc.autofit = False
    
    inc_headers = ["Feature / Capability", "Current Status in Demo", "What Will Happen in Phase 2"]
    inc_widths = [Inches(1.8), Inches(2.3), Inches(2.4)]
    
    for c_idx, (th_text, width) in enumerate(zip(inc_headers, inc_widths)):
        cell = table_inc.cell(0, c_idx)
        cell.width = width
        set_cell_shading(cell, HEX_THEAD)
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(th_text)
        r.bold = True
        r.font.name = 'Calibri'
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    inc_data = [
        (
            "Online Payment Gateway (UPI, Cards & Net Banking)",
            "Currently shows 'Pay at Office' instructions and bank account details. Cash, Cheque, and DD receipts work fully.",
            "Will connect directly to your college bank account via a licensed payment partner (e.g. Razorpay or Cashfree) for instant UPI & QR code payments."
        ),
        (
            "SMS & WhatsApp Parent Alerts",
            "Attendance warnings and fee receipts are currently visible inside the portal login.",
            "Will send automated SMS or WhatsApp messages to parent mobile numbers whenever a student is marked absent or a fee payment is confirmed."
        ),
        (
            "Biometric Fingerprint / RFID Turnstiles",
            "Teachers currently take attendance with one tap on their mobile phone or laptop in the classroom.",
            "Can link directly to electronic biometric or RFID card readers installed at classroom doors or campus gates."
        ),
        (
            "Bus Tracking & Transportation GPS",
            "Transport fee collection is fully working in the fee ledger.",
            "Can be paired with physical GPS trackers installed on college buses to show live vehicle locations on student phones."
        )
    ]

    for r_idx, (f_name, f_cur, f_next) in enumerate(inc_data, start=1):
        bg = HEX_LIGHT_BG if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, (text, width) in enumerate(zip([f_name, f_cur, f_next], inc_widths)):
            cell = table_inc.cell(r_idx, c_idx)
            cell.width = width
            set_cell_shading(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.name = 'Calibri'
            r.font.size = Pt(9)
            if c_idx == 0:
                r.bold = True
                r.font.color.rgb = COLOR_PRIMARY
            else:
                r.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # ----------------------------------------------------
    # 7. ABOUT DOWNTIME & RELIABILITY
    # ----------------------------------------------------
    add_header(doc, "7. About System Reliability & Downtime", level=1)
    
    p_rel = doc.add_paragraph()
    p_rel.paragraph_format.space_after = Pt(8)
    p_rel.paragraph_format.line_spacing = 1.15
    r_rel = p_rel.add_run(
        "We understand that as a college administrator, your biggest concern is whether the system will stop working during peak admission "
        "or exam fee days. The system is hosted on dedicated, professional enterprise cloud infrastructure designed to remain available "
        "around the clock, 365 days a year. Every piece of student data and financial transaction is mirrored continuously and backed up "
        "every single night, ensuring that your records can never be lost even in an emergency. Furthermore, automated monitoring watches "
        "the application 24 hours a day; if any slowdown or hiccup is detected, our technical team is alerted immediately to resolve it "
        "before your staff or students experience any disruption."
    )
    r_rel.font.name = 'Calibri'
    r_rel.font.size = Pt(10.5)
    r_rel.font.color.rgb = COLOR_DARK

    add_callout(
        doc,
        "Zero Lost Data Guarantee: Even if your office internet connection is temporarily interrupted, "
        "the central system remains safe and untouched. When your staff re-connects, all records and balances appear exactly as they were.",
        "Reliability Summary"
    )

    doc.add_page_break()

    # ----------------------------------------------------
    # 8. FEEDBACK NEEDED FROM YOU
    # ----------------------------------------------------
    add_header(doc, "8. Actionable Feedback Questionnaire", level=1)
    
    p_fb_intro = doc.add_paragraph()
    p_fb_intro.paragraph_format.space_after = Pt(10)
    p_fb_intro.paragraph_format.line_spacing = 1.15
    r_fbi = p_fb_intro.add_run(
        "To help us tailor the final system to your exact college guidelines, please review the 7 simple questions below. "
        "You can check the boxes and write quick notes directly so our team knows what adjustments to make."
    )
    r_fbi.font.name = 'Calibri'
    r_fbi.font.size = Pt(10.5)
    r_fbi.font.color.rgb = COLOR_DARK

    questions = [
        (
            "Question 1: Overall Screen Design & Layout",
            "Is the layout of the screens clean, clear, and easy for your office clerks and teachers to read?",
            ["[  ] Yes, the design is clear and easy to navigate", "[  ] Looks good, but some text should be larger", "[  ] Needs simplification (please describe in notes)"]
        ),
        (
            "Question 2: Information on Fee Receipts",
            "Does the official fee receipt contain all the necessary details for your college audits, or should additional fields be added (such as College Affiliation Code, PAN Number, or Bank Account Details)?",
            ["[  ] Current details are sufficient", "[  ] Add Affiliation Number & College Registration Code", "[  ] Add specific bank details for parents"]
        ),
        (
            "Question 3: Attendance Options for Teachers",
            "Is marking attendance by 'Present', 'Absent', and 'Late' sufficient, or do you need special institutional leave categories (such as 'Medical Leave', 'Duty Leave', or 'Suspended')?",
            ["[  ] Present / Absent / Late is sufficient", "[  ] Add Medical Leave & On-Duty options", "[  ] Keep only Present and Absent"]
        ),
        (
            "Question 4: College Fee Particulars",
            "Are the current fee heads (Tuition, Hostel, Exam, Transport, Library, Lab) aligned with your billing rules, or are any specific heads missing?",
            ["[  ] The current fee heads match our college rules", "[  ] Add Building Fund / Development Fee", "[  ] Add Alumni Association or Sports Fee"]
        ),
        (
            "Question 5: Library Loan Duration & Late Fines",
            "Is the default 14-day book loan period with a ₹2 per day overdue fine suitable for your students?",
            ["[  ] Yes, 14 days and ₹2/day is suitable", "[  ] Change loan period to 7 days or 21 days", "[  ] Change fine amount or waive fines completely"]
        ),
        (
            "Question 6: Cashier Daily Shift Closing",
            "At the end of each day, should clerks generate a locked 'Day-End Cash Settlement Report' to submit with physical cash to the Bursar?",
            ["[  ] Yes, clerks must lock and submit daily cash totals", "[  ] No, let cash register remain open for supervisor review"]
        ),
        (
            "Question 7: Approval to Proceed to Final Setup",
            "Based on this working demonstration, are you satisfied with the overall workflow to proceed with finalizing the live system?",
            ["[  ] Yes, proceed with final setup and deployment", "[  ] Make the minor changes noted above first, then proceed", "[  ] Require an in-person walkthrough meeting"]
        )
    ]

    for q_title, q_desc, q_options in questions:
        p_qt = doc.add_paragraph()
        p_qt.paragraph_format.space_before = Pt(8)
        p_qt.paragraph_format.space_after = Pt(2)
        r_qt = p_qt.add_run(q_title)
        r_qt.bold = True
        r_qt.font.name = 'Calibri'
        r_qt.font.size = Pt(11)
        r_qt.font.color.rgb = COLOR_PRIMARY
        
        p_qd = doc.add_paragraph()
        p_qd.paragraph_format.space_after = Pt(4)
        r_qd = p_qd.add_run(q_desc)
        r_qd.font.name = 'Calibri'
        r_qd.font.size = Pt(10)
        r_qd.font.color.rgb = COLOR_DARK
        
        for opt in q_options:
            p_opt = doc.add_paragraph()
            p_opt.paragraph_format.left_indent = Inches(0.25)
            p_opt.paragraph_format.space_after = Pt(2)
            r_opt = p_opt.add_run(opt)
            r_opt.font.name = 'Calibri'
            r_opt.font.size = Pt(9.5)
            r_opt.font.color.rgb = COLOR_MUTED
            
        p_note = doc.add_paragraph()
        p_note.paragraph_format.left_indent = Inches(0.25)
        p_note.paragraph_format.space_after = Pt(8)
        r_note = p_note.add_run("Administrator's Notes: __________________________________________________________________")
        r_note.font.name = 'Calibri'
        r_note.font.size = Pt(9)
        r_note.font.color.rgb = RGBColor(160, 174, 192)

    # Final sign-off block
    p_so = doc.add_paragraph()
    p_so.paragraph_format.space_before = Pt(18)
    p_so.paragraph_format.space_after = Pt(6)
    r_so = p_so.add_run("Reviewed & Acknowledged By:\n\n\n____________________________\t\t\t\t____________________________\nCollege Administrator Signature\t\t\t\tDate")
    r_so.bold = True
    r_so.font.name = 'Calibri'
    r_so.font.size = Pt(10)
    r_so.font.color.rgb = COLOR_PRIMARY

    # Save to both local workspace and artifact directory
    doc.save(OUTPUT_DOCX)
    doc.save(ARTIFACT_DOCX)
    print(f"[SUCCESS] Successfully saved Word document to:\n1. {OUTPUT_DOCX}\n2. {ARTIFACT_DOCX}")

if __name__ == "__main__":
    build_walkthrough()
